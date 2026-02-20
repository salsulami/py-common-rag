"""PowerPoint parser component."""

from __future__ import annotations

from collections.abc import Iterator
from typing import Any

from rag_common.exceptions import MissingDependencyError
from rag_common.parsers.base import BaseDocumentParser
from rag_common.parsers.rendering import iter_office_page_images
from rag_common.prompts import PROMPT_PPTX_CLEANUP, PROMPT_PPTX_VISION_EXTRACTION
from rag_common.types import (
    DocumentChunk,
    FileManifest,
    ParsedDocument,
    VisualExtractionItem,
    VisualExtractionResult,
)


class PptxParser(BaseDocumentParser):
    """Parses PPTX files into slide-oriented chunks."""

    prompt_key = PROMPT_PPTX_CLEANUP
    vision_prompt_key = PROMPT_PPTX_VISION_EXTRACTION

    def parse(self, source_path: str) -> ParsedDocument:
        path = self._resolve_source_path(source_path)
        presentation_cls = _load_presentation()
        presentation = presentation_cls(str(path))

        chunks: list[DocumentChunk] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_lines: list[str] = []
            for shape in slide.shapes:
                slide_lines.extend(_extract_shape_text(shape))

            slide_text = "\n".join(line.strip() for line in slide_lines if line.strip()).strip()
            if not slide_text:
                continue

            chunks.append(
                DocumentChunk(
                    text=slide_text,
                    order=slide_number,
                    metadata={"slide_number": slide_number},
                )
            )

        chunks = self._normalize_chunks(chunks)
        content = "\n\n".join(chunk.text for chunk in chunks).strip()
        return ParsedDocument(
            source_path=str(path),
            content=content,
            chunks=chunks,
            metadata={
                "parser": "pptx",
                "slide_count": len(presentation.slides),
                "chunk_count": len(chunks),
            },
        )

    def iter_visual_items(
        self,
        source_path: str,
        *,
        dpi: int = 170,
    ) -> Iterator[VisualExtractionItem]:
        path = self._resolve_source_path(source_path)
        for slide_index, image_bytes in iter_office_page_images(path, dpi=dpi):
            yield self._extract_visual_item(
                source_path=str(path),
                index=slide_index,
                item_type="slide",
                image_bytes=image_bytes,
                prompt_key=self.vision_prompt_key,
            )

    def extract_visual_result(
        self,
        source_path: str,
        *,
        dpi: int = 170,
        continue_on_error: bool | None = None,
        max_errors: int | None = None,
    ) -> VisualExtractionResult:
        path = self._resolve_source_path(source_path)
        items, errors, attempted = self._collect_visual_items(
            source_path=str(path),
            item_type="slide",
            prompt_key=self.vision_prompt_key,
            images=iter_office_page_images(path, dpi=dpi),
            continue_on_error=continue_on_error,
            max_errors=max_errors,
        )
        manifest = FileManifest(
            source_path=str(path),
            file_type=path.suffix.lower().lstrip("."),
            item_type="slide",
            item_count=len(items),
            metadata={
                "dpi": dpi,
                "attempted_items": attempted,
                "successful_items": len(items),
                "failed_items": len(errors),
            },
        )
        return VisualExtractionResult(file_manifest=manifest, items=items, errors=errors)


def _extract_shape_text(shape: Any) -> list[str]:
    lines: list[str] = []

    has_text_frame = bool(getattr(shape, "has_text_frame", False))
    if has_text_frame:
        text = getattr(shape, "text", "")
        if isinstance(text, str) and text.strip():
            lines.append(text.strip())

    has_table = bool(getattr(shape, "has_table", False))
    if has_table:
        table = getattr(shape, "table", None)
        if table is not None:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if row_cells:
                    lines.append(" | ".join(row_cells))

    return lines


def _load_presentation():
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise MissingDependencyError(
            "PptxParser requires 'python-pptx'. Install with: pip install python-pptx"
        ) from exc
    return Presentation
